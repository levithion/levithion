import sqlite3

import pytest
from openpyxl import load_workbook

from ipv import db, pipeline


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    return pipeline.run(out_dir=tmp_path_factory.mktemp("out"), n_positions=400)


def test_outputs_exist(run):
    out = run.out_dir
    assert (out / "dashboard.html").stat().st_size > 5_000
    assert (out / "ipv.db").exists() and (out / "commentary.md").exists()
    for t in ("dim_position", "dim_date", "fact_ipv", "fact_exception", "fact_control", "fact_desk_bias"):
        assert (out / "powerbi" / f"{t}.csv").exists()


def test_excel_pack_has_named_tables(run):
    wb = load_workbook(next(run.out_dir.glob("IPV_Pack_*.xlsx")))
    assert {"Summary", "Exceptions", "AVA", "Level_Transfers", "ML_Review", "Controls", "Detection"} <= set(wb.sheetnames)
    assert "tblExceptions" in wb["Exceptions"].tables


def test_reconciliation_catches_dropped_record(run):
    assert run.detection_summary["dropped_records_caught_by_recon"] == 1
    assert (run.controls.query("control == 'RECON_COUNT'").status == "FAIL").sum() == 1


def test_detection_floor(run):
    d = run.detection
    assert d.loc["fat_finger", "rules"] == 1.0
    assert d.loc["no_quotes", "rules"] == 1.0
    assert run.detection_summary["rule_false_positive_rate"] < 0.03


def test_numbers_agree_across_outputs(run):
    md = (run.out_dir / "commentary.md").read_text()
    assert run.facts["breach_adj_usd"] in md
    assert run.facts["breach_adj_usd"].lstrip("-") in (run.out_dir / "dashboard.html").read_text()


def test_exception_log_survives_rerun_and_enforces_four_eyes(run):
    con = db.connect(run.out_dir / "ipv.db")
    exc_id = con.execute("SELECT exception_id FROM ipv_exception WHERE status='OPEN' LIMIT 1").fetchone()[0]
    with pytest.raises(sqlite3.IntegrityError):
        db.sign_off(con, exc_id, "ipv-engine", "EXPLAINED", "self-review")   # preparer == reviewer
    db.sign_off(con, exc_id, "analyst.b", "EXPLAINED", "Broker quotes stale; desk mark evidenced by trade.")
    n_before = con.execute("SELECT COUNT(*) FROM ipv_exception").fetchone()[0]
    db.upsert_exceptions(con, run.results)                                   # same month-end rerun
    row = con.execute("SELECT status, commentary FROM ipv_exception WHERE exception_id=?", (exc_id,)).fetchone()
    assert row == ("EXPLAINED", "Broker quotes stale; desk mark evidenced by trade.")
    assert con.execute("SELECT COUNT(*) FROM ipv_exception").fetchone()[0] == n_before
