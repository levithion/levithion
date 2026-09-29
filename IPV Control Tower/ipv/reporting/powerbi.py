"""Star-schema CSV and Excel export for the Power BI model described in powerbi/README.md."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


def export(out: Path, inst: pd.DataFrame, res: pd.DataFrame, exceptions: pd.DataFrame,
           controls: pd.DataFrame, bias: pd.DataFrame) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    dates = pd.DataFrame({"asof": sorted(res["asof"].unique())})
    dates["month"] = dates["asof"].dt.strftime("%b %Y")
    dates["month_index"] = range(1, len(dates) + 1)
    dates["is_latest"] = dates.month_index == len(dates)

    fact = res[["asof", "position_id", "derived_level", "desk_mark", "consensus", "n_quotes", "variance_u",
                "tolerance_u", "tol_ratio", "pv_adjustment_usd", "ava_mpu_usd", "ava_coc_usd", "quote_age_bd",
                "mark_age_periods", "rule_flags", "anomaly_score", "anomaly_flag", "status", "severity"]].copy()
    fact["is_breach"] = fact.rule_flags.str.contains("BREACH|FAT_FINGER")

    tables = {
        "dim_position": inst.drop(columns=["unit_scale"]),
        "dim_date": dates,
        "fact_ipv": fact,
        "fact_exception": exceptions,
        "fact_control": controls,
        "fact_desk_bias": bias,
    }
    paths = []
    for name, df in tables.items():
        p = out / f"{name}.csv"
        df.to_csv(p, index=False, date_format="%Y-%m-%d")
        paths.append(p)
    paths.append(_workbook(out / "IPV_Star_Schema.xlsx", tables))
    return paths


def _workbook(path: Path, tables: dict[str, pd.DataFrame]) -> Path:
    """One workbook, one sheet and one named Excel table per star-schema table, so the Power BI
    service (which imports Excel but not folders of CSVs) discovers all six tables in one upload."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in tables.items():
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([None if pd.isna(v) else v.to_pydatetime() if isinstance(v, pd.Timestamp) else v.item() if hasattr(v, "item") else v
                       for v in row])
        for j, dtype in enumerate(df.dtypes, start=1):
            if pd.api.types.is_datetime64_any_dtype(dtype):
                for cell in ws.iter_cols(min_col=j, max_col=j, min_row=2):
                    for c in cell:
                        c.number_format = "yyyy-mm-dd"
        if df.empty:
            continue
        ref = f"A1:{ws.cell(row=len(df) + 1, column=len(df.columns)).coordinate}"
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(t)
    wb.save(path)
    return path
