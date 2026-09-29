"""Month-end IPV pack as a formatted Excel workbook.

Every sheet is an Excel Table (ListObject) so the VBA toolkit in
excel/IPV_Toolkit.bas and any pivot a reviewer builds can address it by
name rather than by cell range.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

INK = "1F2937"
HEAD = PatternFill("solid", fgColor="1F3A5F")
SEV_FILL = {"CRITICAL": "F8D0D0", "HIGH": "FBE3CF", "MEDIUM": "FFF3C4"}
USD_FMT = '#,##0;[Red]-#,##0'


def _table(ws, df: pd.DataFrame, name: str, start_row: int = 1, money: tuple = (), pct: tuple = (),
           dec: tuple = ()) -> None:
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=start_row, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime().date()
            elif pd.isna(v) if not isinstance(v, (list, dict)) else False:
                v = None
            ws.cell(row=i, column=j, value=v)
    last = start_row + max(len(df), 1)
    ref = f"A{start_row}:{get_column_letter(len(df.columns))}{last}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    for j, c in enumerate(df.columns, 1):
        col = get_column_letter(j)
        fmt = USD_FMT if c in money else "0.0%" if c in pct else "0.00" if c in dec else None
        if fmt:
            for r in range(start_row + 1, last + 1):
                ws[f"{col}{r}"].number_format = fmt
        width = max([len(str(c))] + [len(str(x)) for x in df[c].head(200).tolist()]) + 2
        ws.column_dimensions[col].width = min(max(width, 9), 48)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def build(path: Path, run_id: str, facts: dict, res: pd.DataFrame, exceptions: pd.DataFrame,
          bias: pd.DataFrame, controls: pd.DataFrame, detection: pd.DataFrame, det_summary: dict,
          cfg_rows: list[tuple], commentary_md: str) -> Path:
    cur = res["asof"].max()
    r = res[res["asof"] == cur]
    wb = Workbook()

    # ---------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"IPV Month-End Pack — {facts['asof']}"
    ws["A1"].font = Font(size=16, bold=True, color=INK)
    ws["A2"] = f"Run {run_id} · synthetic data · all adjustments USD · negative = desk mark flatters P&L"
    ws["A2"].font = Font(italic=True, color="6B7280")
    kpis = [("Positions verified", facts["positions"]), ("Independent coverage", facts["coverage_pct"]),
            ("Open exceptions", facts["exceptions"]), ("  of which critical", facts["critical"]),
            ("Proposed IPV adjustment", facts["breach_adj_usd"]), ("Aggregated AVA", facts["ava_usd"]),
            ("Level downgrades / upgrades", f"{facts['level_downgrades']} / {facts['level_upgrades']}"),
            ("ML review queue", facts["ml_review_queue"]), ("Failed data controls", facts["control_failures"])]
    for i, (k, v) in enumerate(kpis, 4):
        ws.cell(row=i, column=1, value=k).font = Font(color="374151")
        ws.cell(row=i, column=2, value=v).font = Font(bold=True, color=INK)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22

    desk = (r.assign(breach=r.pv_adjustment_usd.where(r.rule_flags.str.contains("BREACH|FAT_FINGER"), 0))
              .groupby("desk")
              .agg(positions=("position_id", "size"), exceptions=("status", lambda s: (s == "EXCEPTION").sum()),
                   breach_adjustment_usd=("breach", "sum"), ava_mpu_usd=("ava_mpu_usd", "sum"),
                   ava_coc_usd=("ava_coc_usd", "sum"))
              .reset_index())
    ws.cell(row=15, column=1, value="By desk").font = Font(bold=True, size=12)
    _table(ws, desk, "tblDeskSummary", start_row=16,
           money=("breach_adjustment_usd", "ava_mpu_usd", "ava_coc_usd"))
    ws.freeze_panes = None
    row = 18 + len(desk)
    ws.cell(row=row, column=1, value="Commentary (draft)").font = Font(bold=True, size=12)
    for k, line in enumerate(commentary_md.replace("**", "").replace("### ", "").splitlines(), row + 1):
        ws.cell(row=k, column=1, value=line)

    # ---------------------------------------------------------------- Exceptions
    ws = wb.create_sheet("Exceptions")
    e = r[r.status == "EXCEPTION"].merge(
        exceptions[["position_id", "exception_id", "first_seen", "periods_open", "status"]]
        .rename(columns={"status": "workflow_status"})
        .loc[lambda d: d.workflow_status != "CLOSED"], on="position_id", how="left")
    rank = e.severity.map({"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2})
    e = e.assign(_r=rank, _a=e.pv_adjustment_usd.abs().fillna(-1)).sort_values(["_r", "_a"], ascending=[True, False])
    cols = ["exception_id", "position_id", "desk", "asset_class", "description", "severity", "rule_flags",
            "booked_level", "derived_level", "desk_mark", "consensus", "n_quotes", "variance_u", "unit",
            "tol_ratio", "pv_adjustment_usd", "first_seen", "periods_open", "workflow_status"]
    e = e[cols].assign(reviewer="", commentary="")
    _table(ws, e, "tblExceptions", money=("pv_adjustment_usd",), dec=("variance_u", "tol_ratio"))
    sev_col = get_column_letter(cols.index("severity") + 1)
    rng = f"A2:{get_column_letter(len(e.columns))}{len(e) + 1}"
    for sev, colour in SEV_FILL.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${sev_col}2="{sev}"'],
                                                       fill=PatternFill("solid", fgColor=colour)))

    # ---------------------------------------------------------------- others
    trend = (res.assign(breach=res.pv_adjustment_usd.where(res.rule_flags.str.contains("BREACH|FAT_FINGER"), 0),
                        exc=(res.status == "EXCEPTION").astype(int))
                .groupby(["asof", "desk"]).agg(exceptions=("exc", "sum"), breach_adjustment_usd=("breach", "sum"),
                                               ava_mpu_usd=("ava_mpu_usd", "sum"))
                .reset_index())
    _table(wb.create_sheet("Desk_Trend"), trend, "tblDeskTrend", money=("breach_adjustment_usd", "ava_mpu_usd"))

    ava = (r.groupby(["desk", "asset_class"]).agg(positions=("position_id", "size"), mpu=("ava_mpu_usd", "sum"),
                                                  coc=("ava_coc_usd", "sum")).reset_index())
    ava["ava_after_diversification"] = 0.5 * (ava.mpu + ava.coc)
    _table(wb.create_sheet("AVA"), ava, "tblAVA", money=("mpu", "coc", "ava_after_diversification"))

    lt = r[r.derived_level != r.booked_level][["position_id", "desk", "asset_class", "description", "booked_level",
                                               "derived_level", "n_quotes", "dispersion_u", "tolerance_u"]]
    _table(wb.create_sheet("Level_Transfers"), lt, "tblLevelTransfers", dec=("dispersion_u", "tolerance_u"))

    ml = r[(r.anomaly_flag == 1) & (r.status == "PASS")].sort_values("anomaly_score", ascending=False)
    ml = ml[["position_id", "desk", "asset_class", "description", "anomaly_score", "anomaly_driver", "tol_ratio",
             "n_quotes", "quote_age_bd", "mark_age_periods", "mark_move_gap"]]
    _table(wb.create_sheet("ML_Review"), ml, "tblMLReview", dec=("anomaly_score", "tol_ratio", "mark_move_gap"))

    _table(wb.create_sheet("Desk_Bias"), bias, "tblDeskBias", pct=("share_same_sign",),
           dec=("mean_tol_ratio", "t_stat"))

    ws = wb.create_sheet("Controls")
    _table(ws, controls, "tblControls")
    st = get_column_letter(list(controls.columns).index("status") + 1)
    ws.conditional_formatting.add(f"{st}2:{st}{len(controls) + 1}",
                                  CellIsRule(operator="equal", formula=['"FAIL"'],
                                             fill=PatternFill("solid", fgColor="F8D0D0")))

    ws = wb.create_sheet("Detection")
    _table(ws, detection.reset_index(), "tblDetection",
           pct=("rules", "ml_triage", "desk_bias_test", "any_layer"))
    for i, (k, v) in enumerate(det_summary.items(), len(detection) + 4):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v).number_format = "0.0%" if "rate" in k else "0"

    _table(wb.create_sheet("Config"), pd.DataFrame(cfg_rows, columns=["parameter", "value"]), "tblConfig")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
